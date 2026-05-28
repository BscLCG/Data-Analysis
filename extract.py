import os
from pptx import Presentation
from docx import Document

repo_dir = r"C:\Users\luisc\OneDrive - Universidad Externado de Colombia\2026 I - Data analysis\repo_github\temas"

for folder in os.listdir(repo_dir):
    folder_path = os.path.join(repo_dir, folder)
    if not os.path.isdir(folder_path):
        continue
    
    out_text = []
    
    for filename in os.listdir(folder_path):
        filepath = os.path.join(folder_path, filename)
        if filename.endswith(".pptx") and not filename.startswith("~$"):
            out_text.append(f"\n--- FILE: {filename} ---")
            try:
                prs = Presentation(filepath)
                for i, slide in enumerate(prs.slides):
                    out_text.append(f"Slide {i+1}:")
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            out_text.append(shape.text.strip())
            except Exception as e:
                out_text.append(f"Error reading pptx: {e}")
        elif filename.endswith(".docx") and not filename.startswith("~$"):
            out_text.append(f"\n--- FILE: {filename} ---")
            try:
                doc = Document(filepath)
                for para in doc.paragraphs:
                    if para.text.strip():
                        out_text.append(para.text.strip())
            except Exception as e:
                out_text.append(f"Error reading docx: {e}")
                
    if out_text:
        with open(os.path.join(folder_path, "raw_content.txt"), "w", encoding="utf-8") as f:
            f.write("\n".join(out_text))
